import asyncio
from contextlib import asynccontextmanager
from io import BytesIO
import json

import aiosqlite
from docx import Document
from fastapi import FastAPI, UploadFile
from httpx import ASGITransport, AsyncClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
import pytest

from app.documents import service
from app.chat.audio_input import normalize_chat_attachments
from app.chat.history import _strip_history_attachments
from routes import files


def pdf_bytes(*, blank=False, mixed=False, password=None):
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=500)
    if not blank:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
            DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 40 420 Td (Document fact: blue moon 729.) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
    if mixed:
        writer.add_blank_page(width=400, height=500)
    if password:
        writer.encrypt(password)
    out = BytesIO(); writer.write(out)
    return out.getvalue()


def docx_bytes():
    document = Document()
    document.add_heading('周末计划', 1)
    document.add_paragraph('先散步，然后吃饭。')
    table = document.add_table(rows=2, cols=2)
    for cell, text in zip((c for r in table.rows for c in r.cells), ('时间', '安排', '晚上', '看电影')):
        cell.text = text
    document.add_paragraph('最后一段在表格之后。')
    output = BytesIO(); document.save(output)
    return output.getvalue()


@pytest.fixture
def storage(monkeypatch, tmp_path):
    uploads = tmp_path / 'uploads'; uploads.mkdir()
    monkeypatch.setattr(service, 'UPLOADS_DIR', uploads)
    monkeypatch.setattr(service, 'DOCUMENTS_DIR', tmp_path / 'documents')
    monkeypatch.setattr(files, 'UPLOADS_DIR', uploads)
    db_path = tmp_path / 'test.db'
    @asynccontextmanager
    async def db():
        async with aiosqlite.connect(db_path) as connection:
            await connection.execute('CREATE TABLE IF NOT EXISTS messages (id TEXT, attachments TEXT)')
            yield connection
    monkeypatch.setattr(files, 'get_db', db)
    return tmp_path, db


def upload(name, body):
    return asyncio.run(service.upload_document(UploadFile(filename=name, file=BytesIO(body))))


@pytest.mark.parametrize('name,body,expected', [
    ('随手记.txt', '你好\n第二行'.encode(), '你好\n第二行'),
    ('旧编码.txt', '简体中文测试'.encode('gb18030'), '简体中文测试'),
    ('编码.txt', '另一份文本'.encode('utf-16'), '另一份文本'),
    ('日记.md', b'# Title\n\n**bold**\n\n    indent\n', '# Title\n\n**bold**\n\n    indent\n'),
    ('报告.pdf', pdf_bytes(), 'blue moon 729'),
    ('计划.docx', docx_bytes(), '最后一段在表格之后。'),
], ids=['utf8', 'gb18030', 'utf16', 'markdown', 'pdf', 'docx'])
def test_real_parser_and_cached_source(storage, name, body, expected):
    item = upload(name, body)
    assert item['type'] == 'document' and item['name'] == name
    assert 'text' not in item
    record = service.read_document(item['id'])
    assert expected in record['text']
    assert record['char_count'] == len(record['text'])
    if name.endswith('.docx'):
        assert record['text'].index('先散步') < record['text'].index('| 时间 | 安排 |') < record['text'].index('最后一段')
    trusted = normalize_chat_attachments([{**item, 'name': '伪造', 'url': '/etc/passwd', 'text': '替换正文'}])[0]
    assert trusted == item


@pytest.mark.parametrize('name,body,error', [
    ('空.txt', b'', '空'),
    ('扫描.pdf', pdf_bytes(blank=True), '扫描'),
    ('密码.pdf', pdf_bytes(password='private'), '密码'),
    ('老文档.doc', b'old binary', 'DOCX'),
    ('坏.docx', b'broken', '无法解析'),
    ('二进制.txt', b'hello\x00world', '无法识别'),
    ('长文.txt', ('文' * 60_001).encode(), '60,000'),
    ('大文档.txt', b'a' * (service.MAX_BYTES + 1), '20 MB'),
], ids=['empty', 'scan', 'password', 'legacy-doc', 'corrupt-docx', 'binary', 'too-long', 'too-large'])
def test_invalid_files_are_explained_and_leave_no_upload(storage, name, body, error):
    with pytest.raises(service.InvalidDocument, match=error):
        upload(name, body)
    assert not list(service.UPLOADS_DIR.iterdir())
    assert not service.DOCUMENTS_DIR.exists()


def test_pdf_partial_extraction_is_explicit(storage):
    item = upload('部分图片.pdf', pdf_bytes(mixed=True))
    assert item['page_count'] == 2 and '2' in item['warnings'][0]
    text = service.read_document(item['id'])['text']
    assert '第 1 页' in text and '第 2 页：未提取到文字' in text


@pytest.mark.parametrize('names', [('小栀', '阿澈'), ('青禾', '临川')])
def test_both_provider_payloads_include_actual_file_text_on_followup(storage, monkeypatch, names):
    import ai_providers
    from app.chat import worldbook
    item = upload('原文.md', '# 唯一事实\n\n蓝月编号是729。'.encode())
    history = [{'role': 'user', 'content': '读这份文件', 'attachments': [item]},
               {'role': 'assistant', 'content': '好。', 'attachments': []},
               {'role': 'user', 'content': '编号是多少？', 'attachments': []}]
    original = json.dumps(history, ensure_ascii=False)
    for policy in ('last_user', 'last_message'):
        _strip_history_attachments(history, policy)
        assert history[0]['attachments'] == [item]
    monkeypatch.setattr(worldbook, 'load_worldbook_names', lambda: names)
    for output in (ai_providers.build_multimodal_messages(history), ai_providers.build_gemini_contents(history)):
        encoded = json.dumps(output, ensure_ascii=False)
        assert '蓝月编号是729' in encoded
        assert names[0] in encoded and names[1] in encoded
        assert 'base64' not in encoded and 'inline_data' not in encoded
        assert '用户' not in encoded and '对方' not in encoded and 'AI' not in encoded
    assert json.dumps(history, ensure_ascii=False) == original


def test_document_context_budget_deduplication_and_missing_source(storage, monkeypatch):
    one = upload('第一份.txt', b'first source')
    two = upload('第二份.txt', b'second source')
    monkeypatch.setattr(service, 'MAX_MESSAGE_CHARS', len('second source'))
    messages = [{'role': 'user', 'content': '', 'attachments': [one, one]},
                {'role': 'user', 'content': '', 'attachments': [two]}]
    result = service.document_messages(messages, user_name='青禾', ai_name='临川')
    assert '未加载正文' in result[0]['content'] and result[0]['content'].count('第一份.txt') == 1
    assert 'second source' in result[1]['content']
    service.document_path(two['id']).unlink()
    result = service.document_messages(messages, user_name='青禾', ai_name='临川')
    assert '不可用' in result[1]['content']
    with pytest.raises(service.InvalidDocument):
        normalize_chat_attachments([two])
    with pytest.raises(service.InvalidDocument):
        normalize_chat_attachments([{'type': 'document', 'id': '../../outside'}])


def test_http_upload_preview_discard_and_sent_file_retention(storage):
    app = FastAPI(); app.include_router(files.router)
    async def scenario():
        async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
            response = await client.post('/api/upload', files={'file': ('文稿.txt', '正文完整保留'.encode(), 'application/octet-stream')})
            assert response.status_code == 200
            item = response.json()
            response = await client.get('/api/documents/' + item['id'])
            assert response.json()['text'] == '正文完整保留'
            async with storage[1]() as db:
                await db.execute('INSERT INTO messages VALUES (?,?)', ('sent', json.dumps([item])))
                await db.commit()
            response = await client.delete('/api/documents/' + item['id'])
            assert response.json()['retained']
            async with storage[1]() as db:
                await db.execute('DELETE FROM messages'); await db.commit()
            assert (await client.delete('/api/documents/' + item['id'])).json()['ok']
            assert (await client.get('/api/documents/' + item['id'])).status_code == 404
    asyncio.run(scenario())
    assert not list(service.UPLOADS_DIR.iterdir())
