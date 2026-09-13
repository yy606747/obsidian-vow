"""Real upload/parser/browser roundtrip; temporary data and zero model requests."""
import asyncio
from io import BytesIO
import json
import os
from pathlib import Path
import socket
import sys
import tempfile
import threading
import time

ROOT = Path(__file__).resolve().parents[1]
DATA = tempfile.TemporaryDirectory(prefix='obsidian-documents-browser-')
os.environ['OBSIDIAN_TEST_MODE'] = '1'
os.environ['OBSIDIAN_DATA_DIR'] = DATA.name
sys.path.insert(0, str(ROOT))

from docx import Document
from fastapi import FastAPI, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from playwright.sync_api import sync_playwright
import uvicorn

from app.documents import service
from app.chat.audio_input import normalize_chat_attachments
from app.chat.history import _strip_history_attachments
from database import get_db
from routes.files import router as files_router

OUTPUT = Path('/tmp/obsidian-documents-browser')
messages, payloads, page_errors = [], [], []
app = FastAPI(); app.include_router(files_router)
app.mount('/static', StaticFiles(directory=ROOT / 'static'), name='static')
app.mount('/public', StaticFiles(directory=ROOT.parent / 'public'), name='public')


@app.get('/chat')
async def chat():
    return FileResponse(ROOT / 'static/chat.html')


@app.get('/api/avatar/{kind}')
async def avatar(kind: str):
    if kind == 'info': return {'ai': 0, 'user': 0}
    return FileResponse(ROOT.parent / 'public/optimized' / ('AIIcon.png' if kind == 'ai' else 'UserIcon.png'))


@app.post('/api/conversations/c1/send')
async def send(request: Request):
    body = await request.json()
    n = len(payloads)
    attachments = normalize_chat_attachments(body.get('attachments'))
    item = dict(id=f'user{n}', conv_id='c1', role='user', content=body['content'], attachments=attachments, created_at=time.time())
    messages.append(item)
    history = [dict(m) for m in messages]
    _strip_history_attachments(history, 'last_user')
    payloads.append(service.document_messages(history, user_name='南枝', ai_name='予安'))
    async with get_db() as db:
        await db.execute('INSERT INTO messages VALUES (?,?)', (item['id'], json.dumps(attachments)))
        await db.commit()
    reply = dict(id=f'ai{n}', conv_id='c1', role='assistant', content='**收到文件了。**', attachments=[], created_at=time.time())
    messages.append(reply)
    async def events():
        for event in [dict(type='start', id=reply['id'], user_message_id=item['id']),
                      dict(type='chunk', content=reply['content']), dict(type='message', message=reply)]:
            yield 'data: ' + json.dumps(event, ensure_ascii=False) + '\n\n'
    return StreamingResponse(events(), media_type='text/event-stream')


@app.api_route('/api/{path:path}', methods=['GET', 'POST', 'PUT', 'DELETE'])
async def fixture(path: str):
    if path == 'worldbook': return dict(ai_name='予安', user_name='南枝')
    if path == 'models': return [dict(key='fixture', audio_input=False)]
    if path == 'chat/daily': return dict(enabled=False, count_initialized=True, inherit_count=30, revision=0, timezone='Asia/Shanghai', next_rollover_at=time.time()+86400)
    if path == 'conversations': return [dict(id='c1', title='文件与聊天', model='fixture')]
    if path == 'conversations/c1/messages': return messages
    if path == 'tts/voices': return dict(voices=[])
    if path == 'heart-whispers': return dict(items=[], total=0)
    if path == 'schedule/missed-recent': return dict(count=0, items=[])
    if path.endswith('stream-status'): return dict(streaming=False)
    return {}


async def setup():
    async with get_db() as db:
        await db.execute('CREATE TABLE messages (id TEXT, attachments TEXT)'); await db.commit()


def main():
    OUTPUT.mkdir(exist_ok=True)
    asyncio.run(setup())
    listener = socket.socket(); listener.bind(('127.0.0.1', 0)); port = listener.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [listener]}, daemon=True); thread.start()
    for _ in range(100):
        if server.started: break
        time.sleep(.02)
    checks = []
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            context = browser.new_context(viewport=dict(width=390, height=844), is_mobile=True, service_workers='block')
            context.add_init_script("localStorage.setItem('obsidian_tts_enabled','false')")
            context.route_web_socket('**/ws', lambda ws: ws.on_message(lambda _: ws.send('{"type":"pong"}')))
            page = context.new_page()
            page.on('pageerror', lambda error: page_errors.append(str(error)))
            page.goto(f'http://127.0.0.1:{port}/chat', wait_until='load')
            page.wait_for_function('chatInitialized && !chatSyncPending')
            # Generate a real PDF containing Chinese text through Chromium's print engine.
            print_page = context.new_page()
            print_page.set_content('<meta charset="utf-8"><h1>文件解析测试</h1><p>PDF编号是842。中文正文完整。</p>')
            pdf = print_page.pdf(); print_page.close()
            word = Document(); word.add_heading('周末计划', 1); word.add_paragraph('Word编号是613。')
            table = word.add_table(rows=2, cols=2)
            for cell, value in zip((c for r in table.rows for c in r.cells), ('地点', '时间', '窗边', '晚上')): cell.text = value
            word_data = BytesIO(); word.save(word_data)
            samples = [('随手记.txt', 'TXT编号是729。\n第二段保留。'.encode(), 'TXT编号是729'),
                       ('日记.md', '# **小标题**\n\nMD编号是951。\n<script>window.injected=1</script>'.encode(), 'MD编号是951'),
                       ('周末计划.docx', word_data.getvalue(), 'Word编号是613'),
                       ('文字版.pdf', pdf, 'PDF编号是842')]
            for name, body, fact in samples:
                page.locator('#fileInput').set_input_files(dict(name=name, mimeType='application/octet-stream', buffer=body))
                page.wait_for_selector('.document-preview')
                assert name in page.locator('.document-preview').inner_text()
                page.locator('.document-preview .document-card').click()
                page.wait_for_function("document.querySelector('[data-document-text]')?.textContent.length > 10")
                text = page.locator('[data-document-text]').inner_text()
                assert fact in ''.join(text.split()), (name, text)
                assert page.locator('.document-reader-body script').count() == 0
                assert page.evaluate('window.injected === undefined')
                assert page.evaluate("document.body.classList.contains('vow-covered')")
                page.keyboard.press('Escape')
                assert page.locator('.document-reader[open]').count() == 0
                assert page.evaluate("!document.body.classList.contains('vow-covered')")
                page.locator('#input').fill('读这份文件。')
                page.locator('#sendBtn').click()
                page.wait_for_function('!sending && pendingAttachments.length === 0')
                assert fact in ''.join(json.dumps(payloads[-1], ensure_ascii=False).split())
                checks.append('upload_preview_send_'+name.rsplit('.', 1)[-1])
            page.reload(wait_until='load')
            page.wait_for_function('chatInitialized && !chatSyncPending')
            assert page.locator('#messages .document-card').count() == 4
            page.locator('#messages .document-card').first.click()
            page.wait_for_function("document.querySelector('[data-document-text]').textContent.includes('TXT编号是729')")
            page.locator('[data-document-close]').click()
            page.locator('#input').fill('刚才几份文件分别是什么编号？')
            page.locator('#sendBtn').click(); page.wait_for_function('!sending')
            for _, _, fact in samples: assert fact in ''.join(json.dumps(payloads[-1], ensure_ascii=False).split())
            checks.append('reload_and_followup_retain_all_file_sources')
            page.locator('#fileInput').set_input_files(dict(name='临时.txt', mimeType='text/plain', buffer=b'discard me'))
            page.wait_for_selector('.document-preview')
            removed_id = page.evaluate('pendingAttachments[0].id')
            page.locator('.document-preview [data-preview-remove]').click()
            page.wait_for_function('pendingAttachments.length === 0')
            for _ in range(60):
                if not service.document_path(removed_id).exists(): break
                page.wait_for_timeout(50)
            assert not service.document_path(removed_id).exists()
            page.locator('#fileInput').set_input_files(dict(name='老文件.doc', mimeType='application/msword', buffer=b'old'))
            page.wait_for_function("document.body.innerText.includes('另存为 .docx')")
            assert page.locator('.document-preview').count() == 0
            checks.append('discard_cleans_unsent_upload_and_legacy_doc_explains_conversion')
            for width in (375, 768, 1440):
                page.set_viewport_size(dict(width=width, height=1000))
                page.evaluate("document.getElementById('messages').scrollTop=0")
                assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
                page.screenshot(path=str(OUTPUT / f'cards-{width}.png'))
                page.locator('#messages .document-card').last.click()
                page.wait_for_function("document.querySelector('[data-document-text]').textContent.replaceAll(' ', '').includes('PDF编号是842')")
                assert page.locator('.document-reader').evaluate('el => el.scrollWidth <= el.clientWidth + 1')
                page.screenshot(path=str(OUTPUT / f'reader-{width}.png'))
                page.locator('[data-document-close]').click()
            checks.append('mobile_tablet_desktop_cards_and_reader_fit')
            assert not page_errors, page_errors
            result = dict(passed=checks, page_errors=page_errors, model_calls=0, real_file_formats=4)
            (OUTPUT / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
            print(json.dumps(result, ensure_ascii=False), flush=True)
            browser.close()
    finally:
        server.should_exit = True; thread.join(timeout=10); listener.close(); DATA.cleanup()


if __name__ == '__main__':
    main()
